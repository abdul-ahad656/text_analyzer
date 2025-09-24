from flask import Flask, request, jsonify
from flask_cors import CORS
from werkzeug.utils import secure_filename
import os, fitz, hashlib, pickle, numpy as np, joblib
from docx import Document
from sentence_transformers import SentenceTransformer
from src.pipeline import process_document_pair, clean_text

# Initialize Flask app
app = Flask(__name__)
CORS(app)  # Enable CORS for frontend-backend communication

# === Global model placeholders ===
bert_model = None
svm_model = None

# === Lazy loaders ===
def get_bert_model():
    global bert_model
    if bert_model is None:
        bert_model = SentenceTransformer("models/fine_tuned_MiniLM")
    return bert_model

def get_svm_model():
    global svm_model
    if svm_model is None:
        svm_model = joblib.load("models/svm/svm_ai_detector.pkl")
    return svm_model

# === Utility: Extract raw text from uploaded file ===
def extract_text(file):
    filename = secure_filename(file.filename)
    ext = os.path.splitext(filename)[1].lower()

    try:
        if ext == ".txt":
            return file.read().decode("utf-8", errors="ignore")
        elif ext == ".pdf":
            text = ""
            pdf = fitz.open(stream=file.read(), filetype="pdf")
            for page in pdf:
                text += page.get_text()
            return text
        elif ext == ".docx":
            doc = Document(file)
            return "\n".join(p.text for p in doc.paragraphs)
        else:
            raise ValueError("Unsupported file type")
    except Exception as e:
        raise ValueError(f"Error reading file: {e}")

# === Utility: Convert similarity score into human-readable verdict ===
def interpret_score(score):
    if score > 0.85:
        return "Highly similar"
    elif score > 0.6:
        return "Possibly related"
    return "Likely original"

# === Health check route ===
@app.route("/", methods=["GET"])
def home():
    return "Flask backend is running!"

# === Route: Compare two uploaded documents ===
@app.route("/check", methods=["POST"])
def check_plagiarism():
    file1 = request.files.get("file1")
    file2 = request.files.get("file2")
    if not file1 or not file2:
        return jsonify({"error": "Both files are required"}), 400

    try:
        text1 = clean_text(extract_text(file1))
        text2 = clean_text(extract_text(file2))
        score = process_document_pair(text1, text2)

        return jsonify({
            "similarity_score": float(round(score, 4)),
            "verdict": interpret_score(score)
        })

    except Exception as e:
        return jsonify({"error": str(e)}), 400


# === Route: Compare file + Add to corpus if unique ===
@app.route("/smart-check", methods=["POST"])
def smart_check():
    file = request.files.get("file")
    if not file:
        return jsonify({"error": "File is required"}), 400

    try:
        text = clean_text(extract_text(file))
        filename = secure_filename(file.filename)
        new_embedding = get_bert_model().encode(text)

        corpus_path = "models/embeddings/corpus_embeddings.pkl"
        if os.path.exists(corpus_path):
            with open(corpus_path, "rb") as f:
                filenames, embeddings = pickle.load(f)
        else:
            filenames, embeddings = [], np.empty((0, 384))

        if len(embeddings) > 0:
            scores = np.dot(embeddings, new_embedding) / (
                np.linalg.norm(embeddings, axis=1) * np.linalg.norm(new_embedding)
            )
            max_score = float(np.max(scores))
            most_similar_doc = filenames[np.argmax(scores)]
        else:
            max_score = 0.0
            most_similar_doc = None

        response = {
            "similarity_score": float(round(max_score, 4)),
            "verdict": interpret_score(max_score),
            "most_similar_doc": most_similar_doc,
            "added_to_corpus": False
        }

        if max_score < 0.8:
            hash_name = hashlib.sha1(text.encode()).hexdigest()[:10]
            new_filename = f"{hash_name}_{filename}"
            save_path = os.path.join("data/corpus", new_filename)

            with open(save_path, "w", encoding="utf-8") as f_out:
                f_out.write(text)

            filenames.append(new_filename)
            embeddings = np.vstack([embeddings, new_embedding])

            with open(corpus_path, "wb") as f_out:
                pickle.dump((filenames, embeddings), f_out)

            response["added_to_corpus"] = True

        return jsonify(response)

    except Exception as e:
        return jsonify({"error": str(e)}), 400

# === Route: Detect AI-generated content using SVM ===
@app.route("/detect-ai", methods=["POST"])
def detect_ai():
    file = request.files.get("file")
    if not file:
        return jsonify({"error": "File is required"}), 400

    try:
        text = extract_text(file)
        model = get_svm_model()
        prediction = model.predict([text])[0]
        confidence = model.predict_proba([text])[0][1]

        return jsonify({
            "ai_score": float(round(confidence, 4)),
            "human_score": float(round(1 - confidence, 4)),
            "verdict": "Likely AI-generated" if confidence > 0.7 else "Likely Human-written",
            "model": "SVM (TF-IDF)"
        })

    except Exception as e:
        return jsonify({"error": str(e)}), 400

# === Start the server ===
if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port)
